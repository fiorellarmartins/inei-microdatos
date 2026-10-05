from copy import deepcopy
from pathlib import Path
from unittest.mock import Mock

import pytest
from click.testing import CliRunner
from openpyxl import Workbook, load_workbook

from inei_microdatos.catalog import filter_catalog, load_catalog
from inei_microdatos.cli import cli
from inei_microdatos.download import _collect_module_tasks, _download_one, download_modules, module_download
from inei_microdatos.geography import (
    GeographyUnavailable, _redatam_tree, _resolve_2007, _resolve_2025, select_census_geography,
    subset_rows, subset_workbook, validate_ubigeo,
)
from inei_microdatos.reader import read_catalog_entry, read_module


def geography(year=2025, code="150101"):
    return {"year": year, "code": code, "names": ["Lima", "Lima", "Lima"][:len(code)//2],
            "departments": ["Lima", "Áncash", "Prov. Const. del Callao"],
            "inline_selection": f'DISTRITO {code} "LIMA",', "source_url": "https://example.com"}


@pytest.mark.parametrize("code", [150101, "", "0", "15010", "1501011", "15/101", "150100", "000101", "270101", "15 OR 1"])
def test_reject_invalid_ubigeo_before_network(code, monkeypatch):
    resolver = Mock(side_effect=AssertionError("No network for invalid codes"))
    monkeypatch.setattr("inei_microdatos.geography._resolve", resolver)
    with pytest.raises(ValueError):
        select_census_geography(filter_catalog(load_catalog(), survey="censo"), code)
    resolver.assert_not_called()


def test_scope_all_adapters_without_mutating_catalog(monkeypatch):
    catalog = filter_catalog(load_catalog(), survey="censo")
    before = deepcopy(catalog)
    resolver = Mock(side_effect=lambda y, c: geography(y, c))
    monkeypatch.setattr("inei_microdatos.geography._resolve", resolver)
    selected = select_census_geography(catalog, "150101")
    assert catalog == before
    assert resolver.call_count == 6
    for year, periods in selected[0]["years"].items():
        mod = periods["Unico"]["modules"][0]
        source, code, fmt, ext = module_download(mod, "CSV")
        assert code.endswith("-ubigeo-150101")
        if year in ("1981", "1993", "2005"):
            assert source["data"]["SELECT"] == "SELUSER"
            assert source["data"]["INLINESELECTION"] == 'DISTRITO 150101 "LIMA",'
            assert source["data"].get("WEIGHT") == before[0]["years"][year]["Unico"]["modules"][0]["redatam_query"]["data"].get("WEIGHT")
        elif year == "2007":
            assert "u=150101&" in source
        else:
            assert source["kind"] == "census_geography"
            assert fmt == "XLSX"
    with pytest.raises(ValueError, match="original national catalog"):
        select_census_geography(selected, "150102")


def test_non_census_catalog_rejected(monkeypatch):
    monkeypatch.setattr("inei_microdatos.geography._resolve", Mock(side_effect=AssertionError()))
    with pytest.raises(ValueError, match="requires"):
        select_census_geography(filter_catalog(load_catalog(), survey="endes"), "150101")


def test_custom_layout_cannot_mix_geographies(tmp_path, monkeypatch):
    monkeypatch.setattr("inei_microdatos.geography._resolve", lambda y, c: geography(y, c))
    catalog = filter_catalog(load_catalog(), survey="censo2025")
    first = _collect_module_tasks(select_census_geography(catalog, "150101"), tmp_path, "XLSX", False, "{module_name}.zip")
    second = _collect_module_tasks(select_census_geography(catalog, "150102"), tmp_path, "XLSX", False, "{module_name}.zip")
    assert not {p for _, p in first} & {p for _, p in second}
    assert all("ubigeo-150101" in str(p) for _, p in first)
    assert all("ubigeo-" not in q["source_path"] for q, _ in first)


ROWS = [
    ("CUADRO Nº 1: POBLACIÓN POR DISTRITO", None), ("Departamento, provincia y distrito", "Total"),
    ("PERÚ", 1000), ("ÁNCASH", 100), ("PROVINCIA LIMA", 80), ("DISTRITO LIMA", 70),
    ("DEPARTAMENTO LIMA", 900), ("PROVINCIA LIMA", 800),
    ("DISTRITO LIMA", 300), ("Hombres", 140), ("Mujeres", 160),
    ("DISTRITO ANCÓN", 500), ("PROVINCIA HUARAL", 100), ("DISTRITO LIMA", 90),
    ("Fuente: INEI", None), ("1/ Nota original", None),
]


def test_district_subset_keeps_headers_notes_and_correct_parent_path():
    result = subset_rows(iter(ROWS), geography())
    assert result == ROWS[:2] + ROWS[8:11] + ROWS[-2:]
    assert subset_rows(iter(ROWS), dict(geography(), names=["Lima", "Lima", "Inexistente"])) == []


def test_department_and_province_subtrees():
    assert subset_rows(iter(ROWS), geography(code="15")) == ROWS[:2] + ROWS[6:]
    assert subset_rows(iter(ROWS), geography(code="1501")) == ROWS[:2] + ROWS[7:12] + ROWS[-2:]
    lima_regions = [("CUADRO 1", None), ("LIMA METROPOLITANA 1/", 3), ("REGIÓN LIMA 2/", 4)]
    assert subset_rows(iter(lima_regions), geography(code="15")) == lima_regions


@pytest.mark.parametrize("heading", ["LIMA METROPOLITANA", "Lima Metropolitana 1/"])
def test_lima_implicit_province(heading):
    rows = [("CUADRO 1",), (heading,), ("DISTRITO LIMA", 314840),
            ("DISTRITO ANCÓN", 1), ("REGIÓN LIMA",), ("PROVINCIA HUARAL",), ("DISTRITO LIMA", 9)]
    assert subset_rows(iter(rows), geography()) == [rows[0], rows[2]]
    assert subset_rows(iter(rows), geography(code="1501")) == rows[:4]


def test_lima_province_and_callao_source_headings():
    rows = [("CUADRO 1",), ("DEPARTAMENTO LIMA",), ("PROVINCIA DE LIMA",),
            ("DISTRITO LIMA", 268352), ("PROV. CONSTITUCIONAL DEL CALLAO", 994494),
            ("DISTRITO CALLAO", 451260), ("DISTRITO BELLAVISTA", 1), ("DEPARTAMENTO CUSCO",)]
    assert subset_rows(iter(rows), geography()) == [rows[0], rows[3]]
    geo = dict(geography(), code="070101", names=["Callao", "Callao", "Callao"])
    assert subset_rows(iter(rows), geo) == [rows[0], rows[5]]
    geo.update(code="0701", names=["Callao", "Callao"])
    assert subset_rows(iter(rows), geo) == [rows[0]] + rows[4:7]


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    wb.active.title = "Districts"
    for row in ROWS:
        wb.active.append(row)
    wb.create_sheet("National only").append(["CUADRO 2: TOTAL NACIONAL", 1000])
    wb.create_sheet("Nota de Presentación").append(["Documentation"])
    wb.save(path)
    return path


def test_subset_download_and_read_cache_isolation(workbook, tmp_path, monkeypatch):
    monkeypatch.setattr("inei_microdatos.geography._resolve", lambda y, c: geography(y, c))
    entry = filter_catalog(load_catalog(), survey="censo2025")[0]
    mod = entry["years"]["2025"]["Unico"]["modules"][0]
    dest = tmp_path / "cache"
    dest.mkdir()
    raw = dest / (mod["module_code"] + ".xlsx")
    raw.write_bytes(workbook.read_bytes())
    network = Mock(side_effect=AssertionError("Should reuse national workbook"))
    monkeypatch.setattr("inei_microdatos.download.requests.get", network)
    frames = read_catalog_entry(entry, "2025", dest=dest, ubigeo="150101")
    assert set(frames) == {"Districts"}
    assert frames["Districts"].iat[2, 1] == 300
    assert frames["Districts"].attrs["ubigeo"] == "150101"
    assert raw.read_bytes() == workbook.read_bytes()
    assert len(list(dest.glob("*.xlsx"))) == 2
    network.assert_not_called()


def test_missing_geography_is_unavailable_not_national_fallback(workbook, tmp_path):
    geo = dict(geography(), code="150199", names=["Lima", "Lima", "Missing"])
    dest = tmp_path / "subset.xlsx"
    with pytest.raises(GeographyUnavailable):
        subset_workbook(workbook, dest, geo)
    assert not dest.exists()
    query = {"kind": "census_geography", "url": "https://example.com", "module_code": "source",
             "source_path": str(workbook), "geography": geo}
    assert _download_one(query, dest) == "unavailable"
    assert not dest.exists()


def test_modern_redatam_names_and_full_ubigeo():
    document = '''<form name="SELECTION"><input type="hidden" name="CODES" value="0,1,">
    <table><tr><td><input name="chk18"></td><td>150101 "Lima, Lima, distrito: Lima",</td></tr></table></form>'''
    _, fields, nodes = _redatam_tree(document.encode(), "utf-8")
    assert fields["CODES"] == "0,1,"
    assert nodes["150101"] == ("18", "Lima")


def test_2025_uses_native_lima_region_and_validates_full_code(monkeypatch):
    responses = [
        {"success": True, "data": [{"codRegion": "27", "nombre": "Lima Metropolitana"}]},
        {"success": True, "data": [{"idGeografia": "8150100", "nombre": "Lima"}]},
        {"success": True, "data": [{"idGeografia": "9150101", "nombre": "Lima"}]},
    ]
    get = Mock(side_effect=[Mock(json=Mock(return_value=r)) for r in responses])
    monkeypatch.setattr("inei_microdatos.geography.requests.get", get)
    assert _resolve_2025("150101")["names"] == ["LIMA", "Lima", "Lima"]
    assert get.call_args_list[1].kwargs["params"] == {"codRegion": "27"}
    assert get.call_args_list[2].kwargs["params"] == {"codRegion": "27", "ccpp": "01"}


def test_2007_resolves_each_level_and_rejects_unknown_district(monkeypatch):
    responses = [
        '<datos><dato><cod>01</cod><nom>AMAZONAS</nom></dato></datos>',
        '<datos><dato><cod>0101</cod><nom>CHACHAPOYAS</nom></dato></datos>',
        '<datos><dato><cod>010101</cod><nom>CHACHAPOYAS</nom></dato></datos>',
    ]
    post = Mock(side_effect=[Mock(content=r.encode()) for r in responses])
    monkeypatch.setattr("inei_microdatos.geography.requests.post", post)
    geo = _resolve_2007("010101")
    assert geo["names"] == ["AMAZONAS", "CHACHAPOYAS", "CHACHAPOYAS"]
    assert [c.kwargs["data"]["ubigeo"] for c in post.call_args_list] == ["", "01", "0101"]
    post.side_effect = [Mock(content=r.encode()) for r in responses]
    with pytest.raises(GeographyUnavailable):
        _resolve_2007("010199")


def test_cli_ubigeo_validation_and_dry_run(tmp_path, monkeypatch):
    monkeypatch.setattr("inei_microdatos.geography._resolve", lambda y, c: geography(y, c))
    catalog = str(Path(__file__).parents[1] / "src/inei_microdatos/data/catalog.json")
    args = ["download", "--catalog", catalog, "--survey", "censo2025", "--dest", str(tmp_path), "--dry-run", "--ubigeo"]
    runner = CliRunner()
    result = runner.invoke(cli, args + ["150101"])
    assert result.exit_code == 0, result.output
    assert "ubigeo-150101" in result.output
    result = runner.invoke(cli, args + ["150100"])
    assert result.exit_code != 0
    assert "no 00 components" in result.output
