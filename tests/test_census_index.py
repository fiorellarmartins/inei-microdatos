from copy import deepcopy
from pathlib import Path
from unittest.mock import Mock

import pytest
from click.testing import CliRunner
from openpyxl import Workbook

from inei_microdatos.catalog import load_catalog, filter_catalog
from inei_microdatos.census import CENSUS_LABEL
from inei_microdatos.census_index import workbook_tables
from inei_microdatos.cli import cli
from inei_microdatos.variables import (
    _collect_modules, _index_one_module, _read_index_file, _BUNDLED_INDEX,
    build_index, save_index, search, search_across_years,
)


@pytest.fixture
def census():
    return deepcopy(filter_catalog(load_catalog(), survey="censo"))


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    wb.active.title = "Nota de Presentación"
    wb.active.append(["Documentation, not a variable"])
    wb.create_sheet("Anexo").append(["Boundary changes"])
    sheet = wb.create_sheet("CUADRO Nº 1")
    sheet.append(["CUADRO Nº 1: POBLACIÓN POR SEXO\nY EDAD"])
    sheet.append(["Departamento", "Total", "Hombres"])
    sheet.append(["PERÚ", 12345, 6000])
    wb.save(path)
    return path


def test_legacy_variables_index_without_running_queries(census, tmp_path, monkeypatch):
    network = Mock(side_effect=AssertionError("Metadata indexing must not run REDATAM"))
    monkeypatch.setattr("inei_microdatos.redatam.export_excel", network)
    monkeypatch.setattr("inei_microdatos.download.requests.get", network)
    catalog = filter_catalog(census, year_max=2007)
    entries = build_index(catalog, dest=tmp_path / "index.gz", progress=False)
    assert len(entries) == 293
    assert sum(v["kind"] == "variable" for e in entries for v in e["variables"]) == 190
    assert sum(v["kind"] == "table" for e in entries for v in e["variables"]) == 103
    result = search("PERSONA.SEXO", index=entries, survey="cpv1981", exact=True)[0]
    assert result["kind"] == "variable"
    assert result["data_kind"] == "aggregate_tables"
    assert result["year"] == "1981"
    assert result["source_url"].endswith("censos1981/redatam/")
    assert "n_rows" not in entries[0]  # No invented microdata dimensions.
    network.assert_not_called()


def test_workbook_titles_exclude_data_and_notes(workbook):
    rows = workbook_tables(workbook)
    assert rows == [{"name": "CUADRO Nº 1", "label": "CUADRO Nº 1: POBLACIÓN POR SEXO Y EDAD",
                     "kind": "table", "table": "CUADRO Nº 1"}]


def test_cached_workbook_and_incremental_merge(census, workbook, tmp_path, monkeypatch):
    catalog = filter_catalog(census, survey="censo2025")
    mods = catalog[0]["years"]["2025"]["Unico"]["modules"]
    del mods[1:]
    cache = tmp_path / "data" / "2025" / "Unico"
    cache.mkdir(parents=True)
    (cache / (mods[0]["module_code"] + ".xlsx")).write_bytes(workbook.read_bytes())
    network = Mock(side_effect=AssertionError("Should reuse workbook"))
    monkeypatch.setattr("inei_microdatos.download._download_one", network)
    dest = tmp_path / "index.gz"
    previous = {"survey": "OTHER", "year": "2020", "period": "Anual", "module_code": "old",
                "module_name": "Example", "variables": [{"name": "x", "label": "Previous"}]}
    save_index([previous], dest)
    entries = build_index(catalog, dest=dest, data_dir=tmp_path / "data", progress=False)
    assert entries[0] == previous
    assert len(entries) == 2
    assert entries[1]["variables"][0]["table"] == "CUADRO Nº 1"
    assert build_index(catalog, dest=dest, data_dir=tmp_path / "data", progress=False) == entries
    assert _read_index_file(dest) == entries
    network.assert_not_called()


def test_workbook_download_is_temporary(census, workbook, monkeypatch):
    descriptor = _collect_modules(filter_catalog(census, survey="cpv2017"))[0]
    downloaded = []
    def download(url, path):
        downloaded.append(path)
        path.write_bytes(workbook.read_bytes())
        return "ok"
    monkeypatch.setattr("inei_microdatos.download._download_one", download)
    result = _index_one_module(descriptor)
    assert result["variables"][0]["kind"] == "table"
    assert not downloaded[0].exists()


def test_index_failure_warns_instead_of_silently_skipping(census, tmp_path, monkeypatch):
    catalog = filter_catalog(census, survey="cpv2017")
    monkeypatch.setattr("inei_microdatos.download._download_one", lambda *args: "failed")
    with pytest.warns(RuntimeWarning, match="Could not index CPV2017"):
        assert build_index(catalog, dest=tmp_path / "index.gz", progress=False) == []


def test_search_aliases_years_types_and_track():
    entries = [{"survey": CENSUS_LABEL, "year": year, "period": "Unico", "module_code": year,
                "module_name": "Población", "data_kind": "aggregate_tables", "format": "XLS",
                "variables": [{"name": "SEXO", "label": "Sexo", "kind": kind}]}
               for year, kind in [("1981", "variable"), ("1993", "variable"), ("2017", "table")]]
    assert len(search("sexo", index=entries, survey="cpv")) == 3
    assert search("sexo", index=entries, survey="censo1993")[0]["year"] == "1993"
    assert search("sexo", index=entries, survey="censo1993", year="1981") == []
    assert search("sexo", index=entries, survey="censo", year=2017)[0]["kind"] == "table"
    assert list(search_across_years("SEXO", index=entries, survey="censo")) == ["1981", "1993"]
    assert search("sex", index=entries, exact=True) == []
    assert len(search("sexo", index=entries, module="pob")) == 3
    assert search("sexo", index=entries, module="hogar") == []


def test_bundled_index_covers_every_census_module(census):
    entries = [e for e in _read_index_file(_BUNDLED_INDEX) if e["survey"] == CENSUS_LABEL]
    expected = {m["module_code"] for m in _collect_modules(census)}
    assert len(entries) == len(expected) == 309
    assert {e["module_code"] for e in entries} == expected
    assert all(e["variables"] for e in entries)
    for year in ("1981", "1993", "2005", "2007", "2017", "2025"):
        assert search("sexo", index=entries, survey="censo", year=year)
    assert search("agua", index=entries, survey="cpv2025")
    assert search("PERSONA.SEXO", index=entries, survey="censo1981", exact=True)


def test_cli_identifies_tables_and_source_variables():
    runner = CliRunner()
    result = runner.invoke(cli, ["search", "sexo", "--survey", "censo2025"])
    assert result.exit_code == 0
    assert "[table]" in result.output
    assert "CPV2025" in result.output
    result = runner.invoke(cli, ["search", "PERSONA.SEXO", "--survey", "cpv1981", "--exact"])
    assert result.exit_code == 0
    assert "[REDATAM variable; aggregate query output]" in result.output
