"""Search metadata for census source variables and published aggregate tables."""

from __future__ import annotations

import re
import tempfile
from pathlib import Path


def workbook_tables(path: Path) -> list[dict]:
    """Read titles from the first rows, without indexing observations or totals."""
    from openpyxl import load_workbook

    tables = []
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                titles = []
                for row in sheet.iter_rows(max_row=5, max_col=5, values_only=True):
                    for value in row:
                        if isinstance(value, str):
                            value = " ".join(value.split())
                            if re.match(r"^CUADRO\b", value, re.I):
                                titles.append(value)
                if titles:
                    tables.append({"name": sheet.title, "label": " ".join(dict.fromkeys(titles)),
                                   "kind": "table", "table": sheet.title})
        finally:
            workbook.close()
    if not tables:
        raise ValueError(f"No census table titles found in {path.name}")
    return tables


def index_census_module(mod_info: dict, data_dir=None) -> dict:
    """Use source variable metadata or table titles, never fabricate columns."""
    from inei_microdatos.download import _download_one

    mod = mod_info["census_module"]
    query = mod.get("redatam_query")
    if query:
        variables = [{"name": query["data"]["VARIABLE"], "label": mod["module_name"],
                      "kind": "variable"}]
    elif mod.get("xlsx_url"):
        cached = list(Path(data_dir).rglob(mod["module_code"] + ".xlsx")) if data_dir else []
        if len(cached) > 1:
            raise ValueError(f"Multiple cached workbooks for {mod['module_code']}")
        if cached:
            variables = workbook_tables(cached[0])
        else:
            with tempfile.TemporaryDirectory(prefix="inei-census-index-") as tmp:
                path = Path(tmp) / (mod["module_code"] + ".xlsx")
                status = _download_one(mod["xlsx_url"], path)
                if status not in ("ok", "skipped"):
                    raise OSError(f"Census workbook download failed: {status}")
                variables = workbook_tables(path)
    elif mod.get("xls_url"):
        # The 2007 catalog already contains the official full table title.
        variables = [{"name": mod["module_code"], "label": mod["module_name"],
                      "kind": "table", "table": "tabDetalle"}]
    else:
        raise ValueError("No indexable census metadata")
    return {
        **{key: mod_info[key] for key in ("survey", "category", "year", "period",
                                         "module_code", "module_name", "format")},
        "data_kind": "aggregate_tables", "source_url": mod_info["source_url"],
        "variables": variables,
    }
