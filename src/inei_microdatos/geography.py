"""Year-specific INEI geography and geographic subsets of census tables."""

from __future__ import annotations

from copy import deepcopy
from functools import lru_cache
from itertools import islice, chain
from pathlib import Path
import re
import unicodedata
from urllib.parse import parse_qs, urlencode, urlsplit, urlunsplit
from xml.etree import ElementTree as ET

import requests
from lxml import html

from inei_microdatos.census import CENSUS_VALUE, CENSUS_2007


class GeographyUnavailable(ValueError):
    """The census or table does not contain the requested geographic selection."""


def validate_ubigeo(ubigeo: str) -> str:
    if not isinstance(ubigeo, str) or not re.fullmatch(r"(?:[0-9]{2}|[0-9]{4}|[0-9]{6})", ubigeo):
        raise ValueError("ubigeo must be a string of 2, 4, or 6 digits (department, province, district)")
    if any(ubigeo[i:i + 2] == "00" for i in range(0, len(ubigeo), 2)):
        raise ValueError("Use 2 digits for a department, 4 for a province, or 6 for a district; no 00 components")
    if not 1 <= int(ubigeo[:2]) <= 25:
        raise ValueError("Use INEI ubigeo department codes 01–25, not the portal's special region codes")
    return ubigeo


def _name(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(c for c in value if not unicodedata.combining(c)).upper()
    value = re.sub(r"\s*\d+/\s*$", "", value)
    value = " ".join(value.split()).strip()
    if value in {"PROV. CONST. DEL CALLAO", "PROV. CONSTITUCIONAL DEL CALLAO", "PROVINCIA CONSTITUCIONAL DEL CALLAO"}:
        return "CALLAO"
    if value in {"REGION LIMA", "LIMA METROPOLITANA"}:
        return "LIMA"
    return value


def _redatam_tree(content: bytes, encoding: str):
    root = html.fromstring(content.decode(encoding))
    forms = root.xpath('//form[@name="SELECTION"]')
    if len(forms) != 1:
        raise ValueError("INEI did not return its geographic selection form")
    form = forms[0]
    fields = {n.get("name"): n.get("value", "") for n in form.xpath('.//input[@type="hidden"]')}
    nodes = {}
    for row in form.xpath('.//tr'):
        check = row.xpath('./td/input[starts-with(@name,"chk")]')
        if not check:
            continue
        text = " ".join(row.text_content().split())
        match = re.search(r'\b([0-9]{2}(?:[0-9]{2}){0,2})\s+"([^"]+)"', text)
        if match:
            code, name = match.groups()
            name = re.sub(r"^.*(?:Departamento|Provincia|Distrito):\s*", "", name, flags=re.I)
            nodes[code] = (check[0].get("name")[3:], name)
    return root, fields, nodes


def _resolve_redatam(year: int, code: str) -> dict:
    modern = year == 2017
    base = "CPV2017DI" if modern else f"CPV{year}"
    url = ("https://censos2017.inei.gob.pe/bininei/RpWebStats.exe/Selection" if modern
           else "http://censos1.inei.gob.pe/cgibin/RpWebEngine.exe/Selection")
    encoding = "utf-8-sig" if modern else "cp1252"
    fields = dict(BASE=base, NODE="0", EXPAND="TRUE", CODES="", CURCHECK="XX", CURCHECKVALUE="XX", FILENAME="")
    if modern:
        fields["lang"] = "esp"
    names = []
    departments = []
    with requests.Session() as session:
        def post():
            response = session.post(url, data=fields, timeout=30)
            response.raise_for_status()
            return _redatam_tree(response.content, encoding)
        root, returned, nodes = post()
        fields.update(returned)
        if code[:2] not in nodes:
            fields["NODE"] = "1"  # Expand the country beneath the database root.
            root, returned, nodes = post()
            fields.update(returned)
        departments = [name for key, (_, name) in nodes.items() if len(key) == 2]
        for length in range(2, len(code) + 1, 2):
            prefix = code[:length]
            if prefix not in nodes:
                raise GeographyUnavailable(f"Ubigeo {code} is not available in census {year}")
            node, name = nodes[prefix]
            names.append(name)
            if length < len(code):
                fields.update(NODE=node, CURCHECK="XX", CURCHECKVALUE="XX")
                root, returned, nodes = post()
                fields.update(returned)
        inline = None
        if not modern:
            fields.update(NODE="", CURCHECK=node, CURCHECKVALUE="SEL")
            root, _, _ = post()
            inline = "".join(root.xpath('//textarea[@name="seltext"]/text()')).strip()
            if not re.search(rf"\b{code}\b", inline):
                raise ValueError("INEI did not confirm the requested geographic selection")
    return {"code": code, "year": year, "names": names, "departments": departments,
            "inline_selection": inline, "source_url": url}


def _resolve_2007(code: str) -> dict:
    names, departments = [], []
    for length in range(0, len(code), 2):
        response = requests.post(CENSUS_2007 + "ubigeo.asp", data={"anio": "07", "ubigeo": code[:length]}, timeout=30)
        response.raise_for_status()
        root = ET.fromstring(response.content)
        rows = {row.findtext("cod"): row.findtext("nom") for row in root.findall("dato")}
        if length == 0:
            departments = list(rows.values())
        if code[:length + 2] not in rows:
            raise GeographyUnavailable(f"Ubigeo {code} is not available in census 2007")
        names.append(rows[code[:length + 2]])
    return {"code": code, "year": 2007, "names": names, "departments": departments,
            "source_url": CENSUS_2007 + "ubigeo.asp"}


def _resolve_2025(code: str) -> dict:
    base = "https://censos2025.inei.gob.pe/api/v1/geografia/"
    def get(path, **params):
        response = requests.get(base + path, params=params, timeout=30)
        response.raise_for_status()
        payload = response.json()
        if payload.get("success") is not True or not payload.get("data"):
            raise GeographyUnavailable(f"No official geography available for {code} in 2025")
        return payload["data"]
    deps = get("departamentos")
    departments = [d["nombre"] for d in deps]
    # INEI publishes Lima province as region 27, the other Lima provinces as 26.
    region = ("27" if code[2:4] == "01" else "26") if code[:2] == "15" else code[:2]
    dep = next((d for d in deps if d["codRegion"] == region), None)
    if not dep:
        raise GeographyUnavailable(f"Ubigeo {code} is not available in census 2025")
    names = ["LIMA" if code[:2] == "15" else dep["nombre"]]
    if len(code) >= 4:
        prov = next((p for p in get("provincias", codRegion=region)
                     if p["idGeografia"][1:] == code[:4] + "00"), None)
        if not prov:
            raise GeographyUnavailable(f"Ubigeo {code} is not available in census 2025")
        names.append(prov["nombre"])
    if len(code) == 6:
        district = next((d for d in get("distritos", codRegion=region, ccpp=code[2:4])
                         if d["idGeografia"][1:] == code), None)
        if not district:
            raise GeographyUnavailable(f"Ubigeo {code} is not available in census 2025")
        names.append(district["nombre"])
    return {"code": code, "year": 2025, "names": names, "departments": departments,
            "source_url": base}


@lru_cache(maxsize=128)
def _resolve(year: int, code: str) -> dict:
    for attempt in range(3):
        try:
            if year in (1981, 1993, 2005, 2017):
                return _resolve_redatam(year, code)
            if year == 2007:
                return _resolve_2007(code)
            if year == 2025:
                return _resolve_2025(code)
            raise GeographyUnavailable(f"Census geography is not supported for {year}")
        except requests.RequestException:
            if attempt == 2:
                raise


def select_census_geography(catalog: list[dict], ubigeo: str) -> list[dict]:
    """Return a new catalog scoped to a census-year ubigeo; never change the input."""
    code = validate_ubigeo(ubigeo)
    if not catalog or any(entry.get("value") != CENSUS_VALUE for entry in catalog):
        raise ValueError("ubigeo selection requires a population/housing census catalog; filter survey='censo' first")
    if any(mod.get("geography") for entry in catalog for periods in entry["years"].values()
           for period in periods.values() for mod in period["modules"]):
        raise ValueError("Select geography from the original national catalog, not an already scoped catalog")
    selected = deepcopy(catalog)
    for entry in selected:
        entry["geography"] = {"code": code}
        for year, periods in entry["years"].items():
            geography = deepcopy(_resolve(int(year), code))
            for period in periods.values():
                period["geography"] = geography
                for mod in period["modules"]:
                    mod["geography"] = geography
                    if mod.get("redatam_query"):
                        mod["redatam_query"]["data"].update(SELECT="SELUSER", INLINESELECTION=geography["inline_selection"])
                    elif mod.get("xls_url"):
                        parts = urlsplit(mod["xls_url"])
                        params = parse_qs(parts.query)
                        params["u"] = [code]
                        mod["xls_url"] = urlunsplit(parts._replace(query=urlencode(params, doseq=True)))
                    elif mod.get("xlsx_url"):
                        mod["geography_query"] = {"kind": "census_geography", "url": mod["xlsx_url"],
                                                  "module_code": mod["module_code"], "geography": geography}
    return selected


def _marker(value, departments: set[str]):
    if not isinstance(value, str):
        return None
    value = " ".join(value.split())
    normalized = _name(value)
    if normalized in {"PERU", "TOTAL PERU"}:
        return -1, normalized
    if normalized in departments:
        return 0, normalized
    match = re.fullmatch(r"(DEPARTAMENTO|PROVINCIA|DISTRITO)\s+(?:DE\s+)?(.+)", value, re.I)
    if match:
        return ("DEPARTAMENTO", "PROVINCIA", "DISTRITO").index(match[1].upper()), _name(match[2])
    return None


def subset_rows(rows, geography: dict):
    """Select complete geographic row blocks and retain headers/footnotes."""
    departments = {_name(n) for n in geography["departments"]}
    target = [_name(n) for n in geography["names"]]
    state = [None, None, None]
    headers, selected, notes = [], [], []
    started = False
    for row in rows:
        first = row[0] if row else None
        marker = _marker(first, departments)
        if marker:
            started = True
            level, name = marker
            if level < 0:
                state = [None, None, None]
            elif level == 0 and name == "CALLAO":
                # The constitutional province is both department and province.
                state = ["CALLAO", "CALLAO", None]
            elif level == 0 and re.match(r"^LIMA METROPOLITANA\b", first.strip(), re.I):
                # These workbooks omit a separate province heading for Lima.
                state = ["LIMA", "LIMA", None]
            else:
                state[level:] = [name] + [None] * (2 - level)
        if not started:
            headers.append(row)
        elif isinstance(first, str) and re.match(r"^\s*(?:Fuente\s*:|Nota\s*:|\d+/\s)", first, re.I):
            notes.append(row)
        elif state[:len(target)] == target:
            selected.append(row)
    return headers + selected + notes if selected else []


def subset_workbook(source: Path, dest: Path, geography: dict) -> None:
    """Create a derived XLSX with source values, headers and geographic blocks."""
    from openpyxl import load_workbook, Workbook
    source_book = load_workbook(source, read_only=True, data_only=True)
    result = Workbook(write_only=True)
    try:
        for sheet in source_book:
            iterator = sheet.iter_rows(values_only=True)
            first_rows = list(islice(iterator, 5))
            if not any(isinstance(value, str) and re.match(r"^\s*CUADRO\b", value, re.I)
                       for row in first_rows for value in row):
                continue  # Presentation notes and annexes are not geographic tables.
            rows = subset_rows(chain(first_rows, iterator), geography)
            if not rows:
                continue
            out = result.create_sheet(sheet.title)
            for row in rows:
                out.append(row)
        if not result.sheetnames:
            raise GeographyUnavailable(f"No tables for ubigeo {geography['code']} in {source.name}")
        result.properties.description = (f"Derived geographic subset; census {geography['year']}; "
                                         f"ubigeo {geography['code']}; {' / '.join(geography['names'])}")
        result.save(dest)
    finally:
        source_book.close()
        result.close()
